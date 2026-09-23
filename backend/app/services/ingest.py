"""Read CSV or XLSX uploads and detect which column holds the place name."""

from __future__ import annotations

import csv
import io
from dataclasses import dataclass, field

NAME_HINTS = ["名称", "name", "公园名称", "公园", "地名", "地址名", "查询名称", "关键词", "标题", "项目名称"]
ADNAME_HINTS = ["行政区", "区县", "所在区", "adname", "地区", "行政区划", "城区"]
DISTRICT_SUFFIXES = ("区", "市", "县", "镇", "街")


@dataclass
class Table:
    columns: list[str]
    rows: list[list[str]]


@dataclass
class Detection:
    name_column: str | None
    adname_column: str | None
    confidence: str  # header | heuristic | none
    scores: dict[str, float] = field(default_factory=dict)


def _decode(content: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    return content.decode("utf-8", errors="replace")


def _read_csv(content: bytes, max_rows: int | None) -> Table:
    reader = csv.reader(io.StringIO(_decode(content)))
    raw = list(reader)
    if not raw:
        return Table([], [])
    columns = [c.strip() for c in raw[0]]
    body = raw[1 : (1 + max_rows)] if max_rows else raw[1:]
    rows = [[(cell or "").strip() for cell in row] for row in body]
    return Table(columns, rows)


def _read_xlsx(content: bytes, max_rows: int | None) -> Table:
    from openpyxl import load_workbook

    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    sheet = workbook.worksheets[0]
    iterator = sheet.iter_rows(values_only=True)
    try:
        header = next(iterator)
    except StopIteration:
        return Table([], [])
    columns = ["" if c is None else str(c).strip() for c in header]
    rows: list[list[str]] = []
    for index, raw in enumerate(iterator):
        if max_rows and index >= max_rows:
            break
        rows.append(["" if cell is None else str(cell).strip() for cell in raw])
    workbook.close()
    return Table(columns, rows)


def read_table(filename: str, content: bytes, max_rows: int | None = None) -> Table:
    suffix = (filename or "").lower().rsplit(".", 1)[-1]
    if suffix in {"xlsx", "xlsm"}:
        return _read_xlsx(content, max_rows)
    return _read_csv(content, max_rows)


def _header_score(column: str, hints: list[str]) -> int:
    lowered = column.strip().lower()
    best = 0
    for hint in hints:
        hint_lower = hint.lower()
        if lowered == hint_lower:
            best = max(best, 3)
        elif hint_lower in lowered:
            best = max(best, 2)
    return best


def _column_values(rows: list[list[str]], index: int) -> list[str]:
    out = []
    for row in rows:
        if index < len(row):
            value = row[index].strip()
            if value:
                out.append(value)
    return out


def score_name_column(values: list[str]) -> float:
    """Higher means the column looks more like a list of place names."""
    if not values:
        return 0.0
    total = len(values)
    texty = sum(1 for v in values if not v.replace(".", "", 1).isdigit())
    lengths = [len(v) for v in values]
    avg_len = sum(lengths) / total
    distinct = len(set(values)) / total
    parkish = sum(1 for v in values if "公园" in v or "景区" in v or "自然保护区" in v)
    score = (texty / total) * 2.0
    if 2 <= avg_len <= 30:
        score += 1.0
    score += distinct
    score += (parkish / total) * 1.5
    return round(score, 3)


def score_adname_column(values: list[str]) -> float:
    if not values:
        return 0.0
    total = len(values)
    short = sum(1 for v in values if 1 < len(v) <= 6)
    suffix = sum(1 for v in values if v.endswith(DISTRICT_SUFFIXES))
    distinct = len(set(values))
    if distinct > max(3, total * 0.5):
        return 0.0
    return round((short / total) + (suffix / total), 3)


def detect_columns(table: Table) -> Detection:
    header_name = None
    header_adname = None
    for column in table.columns:
        if header_name is None and _header_score(column, NAME_HINTS):
            header_name = column
        if header_adname is None and _header_score(column, ADNAME_HINTS):
            header_adname = column

    name_scores: dict[str, float] = {}
    adname_scores: dict[str, float] = {}
    for index, column in enumerate(table.columns):
        values = _column_values(table.rows, index)
        name_scores[column] = score_name_column(values)
        adname_scores[column] = score_adname_column(values)

    if header_name:
        return Detection(header_name, header_adname, "header", name_scores)

    best_name = max(name_scores, key=name_scores.get) if name_scores else None
    if best_name and name_scores[best_name] > 1.5:
        candidates = {k: v for k, v in adname_scores.items() if k != best_name}
        best_adname = max(candidates, key=candidates.get) if candidates else None
        if best_adname and adname_scores[best_adname] <= 0.5:
            best_adname = None
        return Detection(best_name, best_adname, "heuristic", name_scores)
    return Detection(None, header_adname, "none", name_scores)
